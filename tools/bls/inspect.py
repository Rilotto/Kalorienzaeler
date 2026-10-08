"""Zeigt den Aufbau des BLS-ZIP: Dateiliste, Spaltenköpfe, erste Zeilen. Nur für die Entwicklung."""
import sys, zipfile, io, csv

path = sys.argv[1]
z = zipfile.ZipFile(path)
print("## Dateien im ZIP")
for i in z.infolist():
    print(f"- {i.filename}  ({i.file_size} Bytes)")

for i in z.infolist():
    n = i.filename.lower()
    data = z.read(i.filename)
    if n.endswith((".csv", ".txt", ".tsv")):
        print(f"\n## {i.filename}")
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc); print(f"Encoding: {enc}"); break
            except UnicodeDecodeError:
                continue
        lines = text.splitlines()
        print(f"Zeilen: {len(lines)}")
        for l in lines[:4]:
            print(l[:3000])
    elif n.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        for ws in wb.worksheets:
            print(f"\n## {i.filename} / Blatt '{ws.title}'  (max_row={ws.max_row}, max_col={ws.max_column})")
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                cells = ["" if v is None else str(v) for v in row]
                print(f"Zeile {r+1}: " + " | ".join(cells)[:4000])
                if r >= 4:
                    break
