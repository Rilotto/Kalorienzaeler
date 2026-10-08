"""Zeigt den Aufbau des BLS-ZIP: Dateiliste, Spaltenköpfe, erste Zeilen. Nur für die Entwicklung."""
import sys, zipfile, io, csv

path = sys.argv[1]
z = zipfile.ZipFile(path)
print("## Dateien im ZIP")
for i in z.infolist():
    print(f"- {i.filename}  ({i.file_size} Bytes)")

for i in z.infolist():
    n = i.filename.lower()
    data = z.read(i.filename)
    if n.endswith((".csv", ".txt", ".tsv")):
        print(f"\n## {i.filename}")
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                text = data.decode(enc); print(f"Encoding: {enc}"); break
            except UnicodeDecodeError:
                continue
        lines = text.splitlines()
        print(f"Zeilen: {len(lines)}")
        for l in lines[:4]:
            print(l[:3000])
    elif n.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
        for ws in wb.worksheets:
            print(f"\n## {i.filename} / Blatt '{ws.title}'  (max_row={ws.max_row}, max_col={ws.max_column})")
            limit = 60 if "Components" in i.filename else 3
            for r, row in enumerate(ws.iter_rows(values_only=True)):
                cells = ["" if v is None else str(v) for v in row]
                if "Components" in i.filename:
                    print(f"Zeile {r+1}: " + " | ".join(cells)[:400])
                else:
                    # nur die ersten 12 Spalten und Spalten, deren Kopf ENERC/PROT/FAT/CHO/WATER enthält
                    if r == 0:
                        header = cells
                        print("Kopf (alle):", " | ".join(cells)[:6000])
                        keep = [k for k, h in enumerate(cells) if k < 12 or any(x in h.upper() for x in ("ENERC", "PROT", "FAT", "CHO", "WATER"))]
                    print(f"Zeile {r+1}: " + " | ".join(f"{header[k] if k < len(header) else k}={cells[k] if k < len(cells) else ''}" for k in keep)[:3000])
                if r >= limit:
                    break
